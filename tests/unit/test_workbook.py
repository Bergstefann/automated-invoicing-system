"""Tests for `invoicing.workbook`, the `.xlsx` loader/validator.

Fixture workbooks are built in-memory with openpyxl rather than committed
as binary files — every malformed case is then reviewable as ordinary
Python in this file's diff, and there's no risk of a stray binary fixture
drifting from what the test actually claims about it. The one *real*,
committed workbook is `templates/lesson_schedule.xlsx` (the Phase 2
deliverable itself); `test_the_committed_template_loads_and_validates`
below is the thing that keeps that file honest.
"""

from __future__ import annotations

import io
from datetime import date, time
from pathlib import Path

import pytest
from openpyxl import Workbook

from invoicing.schedule_contract import ScheduleContractError
from invoicing.workbook import load_schedule_workbook, to_schedule_snapshot

STUDENT_HEADERS = ["student_id", "name", "rate_cents", "parent_name", "parent_email"]
SCHEDULE_HEADERS = ["student_id", "lesson_date", "start_time", "duration_minutes", "status"]

DEFAULT_STUDENTS = [
    ["S-0001", "Aria Smith", 4000, "Smith Parent", "smith@example.com"],
    ["S-0002", "Dax Jones", 4500, "Jones Parent", "jones@example.com"],
]
DEFAULT_SCHEDULE = [
    ["S-0001", date(2026, 3, 2), time(16, 0), 30, "Y"],
    ["S-0002", date(2026, 3, 2), time(16, 30), 30, "YI"],
]


def _build_workbook(
    *,
    schema_version: object = 1,
    student_headers: list[str] | None = None,
    student_rows: list[list[object]] | None = None,
    schedule_headers: list[str] | None = None,
    schedule_rows: list[list[object]] | None = None,
    include_students_sheet: bool = True,
    include_schedule_sheet: bool = True,
) -> io.BytesIO:
    wb = Workbook()
    legend = wb.active
    legend.title = "Legend"
    legend["A2"] = "Schema Version"
    legend["B2"] = schema_version

    if include_students_sheet:
        students = wb.create_sheet("Students")
        students.append(student_headers if student_headers is not None else STUDENT_HEADERS)
        for row in DEFAULT_STUDENTS if student_rows is None else student_rows:
            students.append(row)

    if include_schedule_sheet:
        schedule = wb.create_sheet("Schedule")
        schedule.append(schedule_headers if schedule_headers is not None else SCHEDULE_HEADERS)
        for row in DEFAULT_SCHEDULE if schedule_rows is None else schedule_rows:
            schedule.append(row)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


# ── valid workbook ──────────────────────────────────────────────────────


def test_loads_a_valid_workbook() -> None:
    result = load_schedule_workbook(_build_workbook())
    assert result.schema_version == 1
    assert set(result.register) == {"S-0001", "S-0002"}
    assert len(result.lessons) == 2


def test_to_schedule_snapshot_carries_student_id_through() -> None:
    result = load_schedule_workbook(_build_workbook())
    snapshot = to_schedule_snapshot(result)
    ids = {s.student_id for s in snapshot.students}
    assert ids == {"S-0001", "S-0002"}
    lesson_ids = {lesson.student_id for lesson in snapshot.lessons}
    assert lesson_ids == {"S-0001", "S-0002"}


def test_column_order_in_the_sheet_does_not_matter() -> None:
    reordered_headers = ["parent_email", "student_id", "rate_cents", "name", "parent_name"]
    reordered_rows = [
        ["smith@example.com", "S-0001", 4000, "Aria Smith", "Smith Parent"],
    ]
    result = load_schedule_workbook(
        _build_workbook(
            student_headers=reordered_headers,
            student_rows=reordered_rows,
            schedule_rows=[DEFAULT_SCHEDULE[0]],
        )
    )
    assert result.register["S-0001"].name == "Aria Smith"


# ── malformed: schema version ───────────────────────────────────────────


def test_rejects_an_unsupported_schema_version() -> None:
    with pytest.raises(ScheduleContractError, match="unsupported schema_version 99"):
        load_schedule_workbook(_build_workbook(schema_version=99))


def test_rejects_a_missing_schema_version() -> None:
    with pytest.raises(ScheduleContractError, match="must be an integer"):
        load_schedule_workbook(_build_workbook(schema_version=None))


# ── malformed: missing sheet / column ───────────────────────────────────


def test_rejects_a_workbook_missing_the_students_sheet() -> None:
    with pytest.raises(ScheduleContractError, match="missing required sheet 'Students'"):
        load_schedule_workbook(_build_workbook(include_students_sheet=False))


def test_rejects_a_workbook_missing_the_schedule_sheet() -> None:
    with pytest.raises(ScheduleContractError, match="missing required sheet 'Schedule'"):
        load_schedule_workbook(_build_workbook(include_schedule_sheet=False))


def test_rejects_a_schedule_sheet_missing_a_required_column() -> None:
    headers_missing_duration = ["student_id", "lesson_date", "start_time", "status"]
    rows_missing_duration = [["S-0001", date(2026, 3, 2), time(16, 0), "Y"]]
    with pytest.raises(
        ScheduleContractError,
        match=r"'Schedule' sheet is missing required column\(s\): \['duration_minutes'\]",
    ):
        load_schedule_workbook(
            _build_workbook(
                schedule_headers=headers_missing_duration,
                schedule_rows=rows_missing_duration,
            )
        )


# ── malformed: register / schedule rows ─────────────────────────────────


def test_rejects_an_unknown_student_id_in_the_schedule() -> None:
    with pytest.raises(ScheduleContractError, match="student_id 'S-0031' not found in register"):
        load_schedule_workbook(
            _build_workbook(schedule_rows=[["S-0031", date(2026, 3, 2), time(16, 0), 30, "Y"]])
        )


def test_rejects_a_duplicate_student_id_in_the_register() -> None:
    duplicated = [
        ["S-0001", "Aria Smith", 4000, "Smith Parent", "smith@example.com"],
        ["S-0001", "Someone Else", 4000, "Other Parent", "other@example.com"],
    ]
    with pytest.raises(ScheduleContractError, match="duplicate student_id 'S-0001'"):
        load_schedule_workbook(_build_workbook(student_rows=duplicated))


def test_rejects_a_bad_date_format_in_the_schedule() -> None:
    with pytest.raises(ScheduleContractError, match="lesson_date must be a date"):
        load_schedule_workbook(
            _build_workbook(schedule_rows=[["S-0001", "2/3/2026", time(16, 0), 30, "Y"]])
        )


def test_attribute_change_on_an_existing_student_id_still_resolves_to_one_student() -> None:
    """Directly echoes the incident: a changed attribute (parent_email)
    must never fork identity — the schedule row still joins to the same
    student_id, and there is still exactly one entry in the register."""
    changed_contact_students = [
        ["S-0001", "Aria Smith", 4000, "Smith Parent", "smith.new-address@example.com"],
    ]
    result = load_schedule_workbook(
        _build_workbook(
            student_rows=changed_contact_students,
            schedule_rows=[DEFAULT_SCHEDULE[0]],
        )
    )
    assert len(result.register) == 1
    assert result.register["S-0001"].parent_email == "smith.new-address@example.com"
    assert len(result.lessons) == 1
    assert result.lessons[0].student_id == "S-0001"


# ── the real, committed template ────────────────────────────────────────


def test_the_committed_template_loads_and_validates() -> None:
    template = Path(__file__).parents[2] / "templates" / "lesson_schedule.xlsx"
    result = load_schedule_workbook(template)
    assert result.schema_version == 1
    assert len(result.register) > 0
    assert len(result.lessons) > 0
