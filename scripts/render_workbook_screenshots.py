"""Renders PNG screenshots of `templates/lesson_schedule.xlsx` for the
README — nobody opens an `.xlsx` from a GitHub file listing, so the
Legend/Students/Schedule sheets get rendered as images from the workbook's
own actual cell values, fills, and fonts (read via openpyxl), the same way
every other portfolio repo in this set does it.

This isn't a real Excel/LibreOffice screenshot (neither is available in
this environment) — it's a from-scratch table renderer over openpyxl's
cell data, sized from the workbook's own column widths. Run directly to
regenerate after the template changes:

    .venv/Scripts/python.exe scripts/render_workbook_screenshots.py
"""

from __future__ import annotations

import textwrap
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image, ImageDraw, ImageFont

WORKBOOK_PATH = Path(__file__).resolve().parents[1] / "templates" / "lesson_schedule.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "docs" / "images"

FONT_DIR = Path("C:/Windows/Fonts")
FONT_REGULAR = FONT_DIR / "calibri.ttf"
FONT_BOLD = FONT_DIR / "calibrib.ttf"
FONT_ITALIC = FONT_DIR / "calibrii.ttf"

CHAR_PX = 7.3  # approx pixels per Excel "character width" unit, at the font size used
ROW_HEIGHT = 26
HEADER_ROW_HEIGHT = 30
ROW_NUM_COL_WIDTH = 36
PADDING = 6
GRID_COLOR = (208, 208, 208)
ROW_HEADER_FILL = (243, 243, 243)
ROW_HEADER_TEXT = (110, 110, 110)
DEFAULT_TEXT = (30, 30, 30)


def _font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(path), size)


def _col_pixel_width(sheet: Worksheet, col_index: int) -> int:
    letter = sheet.cell(row=1, column=col_index).column_letter
    dim = sheet.column_dimensions.get(letter)
    width_units = dim.width if dim and dim.width else 10.0
    return max(60, int(width_units * CHAR_PX))


def _cell_text_color(cell: object) -> tuple[int, int, int]:
    font = cell.font  # type: ignore[attr-defined]
    color = font.color
    if color is not None and color.type == "rgb" and isinstance(color.rgb, str):
        rgb = color.rgb[-6:]
        if rgb != "000000":
            return tuple(int(rgb[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]
    return DEFAULT_TEXT


def _cell_fill_color(cell: object) -> tuple[int, int, int] | None:
    fill = cell.fill  # type: ignore[attr-defined]
    if fill is None or fill.fgColor is None or fill.fgColor.type != "rgb":
        return None
    rgb = fill.fgColor.rgb
    if not isinstance(rgb, str) or rgb in ("00000000",):
        return None
    hex6 = rgb[-6:]
    return tuple(int(hex6[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


def _format_value(cell: object) -> str:
    value = cell.value  # type: ignore[attr-defined]
    if value is None:
        return ""
    fmt = (cell.number_format or "").upper()  # type: ignore[attr-defined]
    if "YYYY-MM-DD" in fmt and hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    if "HH:MM" in fmt and hasattr(value, "strftime"):
        return value.strftime("%H:%M")
    if fmt == "@":  # text-formatted cell
        return str(value)
    return str(value)


def render_sheet(sheet: Worksheet, out_path: Path, *, max_row: int, max_col: int) -> None:
    col_widths = [ROW_NUM_COL_WIDTH] + [_col_pixel_width(sheet, c) for c in range(1, max_col + 1)]

    row_heights = []
    for r in range(1, max_row + 1):
        tallest = HEADER_ROW_HEIGHT if r == 1 else ROW_HEIGHT
        for c in range(1, max_col + 1):
            cell = sheet.cell(row=r, column=c)
            text = _format_value(cell)
            if not text:
                continue
            wrap_width = max(10, int((col_widths[c] - 2 * PADDING) / (CHAR_PX * 0.92)))
            lines = textwrap.wrap(text, width=wrap_width) or [""]
            needed = len(lines) * 18 + 2 * PADDING
            tallest = max(tallest, needed)
        row_heights.append(tallest)

    col_header_height = 22
    img_width = sum(col_widths) + 2
    img_height = col_header_height + sum(row_heights) + 2
    img = Image.new("RGB", (img_width, img_height), "white")
    draw = ImageDraw.Draw(img)

    small_font = _font(FONT_REGULAR, 11)

    # column-letter header strip
    x = col_widths[0]
    draw.rectangle([0, 0, img_width, col_header_height], fill=ROW_HEADER_FILL)
    for c in range(1, max_col + 1):
        letter = sheet.cell(row=1, column=c).column_letter
        w = col_widths[c]
        draw.text(
            (x + w / 2, col_header_height / 2),
            letter,
            font=small_font,
            fill=ROW_HEADER_TEXT,
            anchor="mm",
        )
        x += w

    y = col_header_height
    for r_index, r in enumerate(range(1, max_row + 1), start=0):
        h = row_heights[r_index]
        # row-number cell
        draw.rectangle([0, y, col_widths[0], y + h], fill=ROW_HEADER_FILL, outline=GRID_COLOR)
        draw.text(
            (col_widths[0] / 2, y + h / 2),
            str(r),
            font=small_font,
            fill=ROW_HEADER_TEXT,
            anchor="mm",
        )

        x = col_widths[0]
        for c in range(1, max_col + 1):
            w = col_widths[c]
            cell = sheet.cell(row=r, column=c)
            fill = _cell_fill_color(cell)
            draw.rectangle([x, y, x + w, y + h], fill=fill or "white", outline=GRID_COLOR)

            text = _format_value(cell)
            if text:
                bold = bool(cell.font.bold)
                italic = bool(cell.font.italic)
                font = _font(FONT_BOLD if bold else (FONT_ITALIC if italic else FONT_REGULAR), 13)
                color = _cell_fill_color(cell) and (255, 255, 255) or _cell_text_color(cell)
                wrap_width = max(10, int((w - 2 * PADDING) / (CHAR_PX * 0.92)))
                lines = textwrap.wrap(text, width=wrap_width) or [""]
                align = (cell.alignment.horizontal or "left") if cell.alignment else "left"
                line_y = y + PADDING
                for line in lines:
                    if align == "center":
                        draw.text((x + w / 2, line_y), line, font=font, fill=color, anchor="ma")
                    else:
                        draw.text((x + PADDING, line_y), line, font=font, fill=color)
                    line_y += 18
            x += w
        y += h

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    img.save(out_path)


def build() -> list[Path]:
    wb = load_workbook(WORKBOOK_PATH)
    outputs = []

    legend = wb["Legend"]
    out = OUTPUT_DIR / "schedule-workbook-legend.png"
    render_sheet(legend, out, max_row=18, max_col=3)
    outputs.append(out)

    students = wb["Students"]
    out = OUTPUT_DIR / "schedule-workbook-students.png"
    render_sheet(students, out, max_row=3, max_col=5)
    outputs.append(out)

    schedule = wb["Schedule"]
    out = OUTPUT_DIR / "schedule-workbook-schedule.png"
    render_sheet(schedule, out, max_row=3, max_col=5)
    outputs.append(out)

    return outputs


if __name__ == "__main__":
    for path in build():
        print(f"Wrote {path}")
