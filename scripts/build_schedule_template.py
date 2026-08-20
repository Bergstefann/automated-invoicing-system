"""Generates `templates/lesson_schedule.xlsx` — the schema-contract
workbook described in `docs/SCHEDULE-SCHEMA.md`.

Run directly to regenerate the template after a contract change:

    .venv/Scripts/python.exe scripts/build_schedule_template.py

The column names, sheet names, and version cell written here must match
`invoicing.schedule_contract` and `invoicing.workbook` exactly — there is a
test (`test_the_committed_template_loads_and_validates` in
tests/unit/test_workbook.py) that loads the committed output of this script
through the real loader, so a drift between the two fails CI rather than
surfacing only when someone opens the file by hand.

Example data is entirely fabricated (fictional names, @example.com
addresses) — nothing here is derived from, or resembles, data belonging to
the real business this pipeline was built for.
"""

from __future__ import annotations

from datetime import date, time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from invoicing.schedule_contract import SUPPORTED_SCHEMA_VERSIONS
from invoicing.workbook import (
    LEGEND_SHEET,
    SCHEDULE_SHEET,
    SCHEMA_VERSION_CELL,
    STUDENTS_SHEET,
)

OUTPUT_PATH = Path(__file__).resolve().parents[1] / "templates" / "lesson_schedule.xlsx"
CURRENT_SCHEMA_VERSION = min(SUPPORTED_SCHEMA_VERSIONS)

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
NOTE_FONT = Font(italic=True, color="555555")

STUDENT_HEADERS = ["student_id", "name", "rate_cents", "parent_name", "parent_email"]
SCHEDULE_HEADERS = ["student_id", "lesson_date", "start_time", "duration_minutes", "status"]

EXAMPLE_STUDENTS = [
    ("S-0001", "Aria Ashworth", 4000, "Jordan Ashworth", "jordan.ashworth@example.com"),
    ("S-0002", "Dax Brambleton", 4500, "Sam Brambleton", "sam.brambleton@example.com"),
]
EXAMPLE_SCHEDULE = [
    ("S-0001", date(2026, 3, 2), time(16, 0), 30, "Y"),
    ("S-0002", date(2026, 3, 2), time(16, 30), 30, "YI"),
]


def _style_header_row(sheet: Worksheet, headers: list[str]) -> None:
    for col_index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet: Worksheet, widths: list[int]) -> None:
    for col_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=col_index).column_letter].width = width


def _build_legend_sheet(wb: Workbook) -> None:
    legend = wb.active
    legend.title = LEGEND_SHEET

    legend["A1"] = "Lesson Schedule — Schema Contract"
    legend["A1"].font = TITLE_FONT

    legend["A2"] = "Schema Version"
    legend["A2"].font = Font(bold=True)
    legend[SCHEMA_VERSION_CELL] = CURRENT_SCHEMA_VERSION
    legend[SCHEMA_VERSION_CELL].alignment = Alignment(horizontal="center")
    legend["C2"] = (
        "Read by the parser to refuse a workbook it doesn't understand. "
        "Don't edit unless you are deliberately migrating to a new contract version."
    )
    legend["C2"].font = NOTE_FONT

    legend["A4"] = "What's editable"
    legend["A4"].font = Font(bold=True)
    editable_rows = [
        "Students sheet: every row below the header. Add a row per new student; "
        "never delete or renumber an existing student_id.",
        "Schedule sheet: every row below the header. One row per lesson.",
        "Legend sheet: nothing, except Schema Version when deliberately migrating.",
    ]
    for offset, text in enumerate(editable_rows):
        legend.cell(row=5 + offset, column=1, value=f"• {text}")

    legend["A9"] = "Column reference"
    legend["A9"].font = Font(bold=True)
    glossary = [
        ("Students.student_id", "Identity. Text, format \"S-0001\". Issued once, "
         "the first time a student appears. Never reused, never edited afterward."),
        ("Students.name / rate_cents / parent_name / parent_email", "Attributes, not "
         "identity — these can change freely across edits without ever creating a "
         "second student. rate_cents is whole cents (4000 = $40.00)."),
        ("Schedule.student_id", "Must match a student_id already in the Students "
         "sheet — pick it from the dropdown, don't type it freehand."),
        ("Schedule.lesson_date / start_time", "Real date / time values, not text."),
        ("Schedule.duration_minutes", "Positive whole number of minutes."),
        ("Schedule.status", "Y = billable, not yet invoiced. YI = already invoiced. "
         "N = not billable (e.g. cancelled)."),
    ]
    for offset, (column, meaning) in enumerate(glossary):
        row = 10 + offset
        legend.cell(row=row, column=1, value=column).font = Font(bold=True)
        legend.cell(row=row, column=2, value=meaning)

    legend["A17"] = "Example row (Schedule sheet)"
    legend["A17"].font = Font(bold=True)
    legend.cell(row=18, column=1, value="S-0001 | 2026-03-02 | 16:00 | 30 | Y")
    legend["A18"].font = NOTE_FONT

    legend.column_dimensions["A"].width = 42
    legend.column_dimensions["B"].width = 70
    legend.column_dimensions["C"].width = 48


def _build_students_sheet(wb: Workbook) -> Worksheet:
    sheet = wb.create_sheet(STUDENTS_SHEET)
    _style_header_row(sheet, STUDENT_HEADERS)
    sheet["A1"].comment = Comment(
        "The identity register. student_id is the only thing that identifies a "
        "student — name, rate, and contact are attributes, and can change freely "
        "without ever forking a second student. See docs/SCHEDULE-SCHEMA.md.",
        "schema-contract",
    )

    for row_offset, row in enumerate(EXAMPLE_STUDENTS, start=2):
        for col_offset, value in enumerate(row, start=1):
            sheet.cell(row=row_offset, column=col_offset, value=value)

    # student_id as text so Excel never strips a leading "S-000" down to a
    # number or mangles it into scientific notation.
    for row_index in range(2, 2 + len(EXAMPLE_STUDENTS) + 200):
        sheet.cell(row=row_index, column=1).number_format = "@"

    _autosize(sheet, [12, 22, 12, 20, 30])
    return sheet


def _build_schedule_sheet(wb: Workbook, students_sheet: Worksheet) -> None:
    sheet = wb.create_sheet(SCHEDULE_SHEET)
    _style_header_row(sheet, SCHEDULE_HEADERS)

    for row_offset, row in enumerate(EXAMPLE_SCHEDULE, start=2):
        for col_offset, value in enumerate(row, start=1):
            sheet.cell(row=row_offset, column=col_offset, value=value)

    max_row = 2 + len(EXAMPLE_SCHEDULE) + 200

    for row_index in range(2, max_row):
        sheet.cell(row=row_index, column=1).number_format = "@"
        sheet.cell(row=row_index, column=2).number_format = "YYYY-MM-DD"
        sheet.cell(row=row_index, column=3).number_format = "HH:MM"

    student_id_validation = DataValidation(
        type="list",
        formula1=f"={students_sheet.title}!$A$2:$A$201",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Unknown student_id",
        error="Pick a student_id from the Students sheet — free text isn't validated.",
    )
    sheet.add_data_validation(student_id_validation)
    student_id_validation.add(f"A2:A{max_row}")

    status_validation = DataValidation(
        type="list",
        formula1='"Y,YI,N"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Unknown status",
        error="Status must be one of Y (billable), YI (invoiced), N (not billable).",
    )
    sheet.add_data_validation(status_validation)
    status_validation.add(f"E2:E{max_row}")

    _autosize(sheet, [12, 14, 11, 20, 10])


def build() -> Path:
    wb = Workbook()
    _build_legend_sheet(wb)
    students_sheet = _build_students_sheet(wb)
    _build_schedule_sheet(wb, students_sheet)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
